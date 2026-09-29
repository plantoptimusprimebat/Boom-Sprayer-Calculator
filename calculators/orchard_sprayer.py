"""Orchard Sprayer timed-catch calibration calculator."""
from __future__ import annotations

import math
import statistics
from datetime import datetime, timedelta, timezone
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
import streamlit as st


if st.button("← Back to calculator selection", type="secondary"):
    st.session_state["selected_calculator"] = None
    st.rerun()


BASE_KEYS = ("orchard_duration", "orchard_unit")
REPLICATIONS = (1, 2, 3)
UNITS = ("ml", "L")


def valid(value: object) -> bool:
    """Return true only for finite numeric values."""
    return value is not None and isinstance(value, (int, float)) and math.isfinite(value)


def show(value: object, decimals: int = 3) -> str:
    """Format a number for display or return a dash when unavailable."""
    return f"{value:,.{decimals}f}" if valid(value) else "—"


def reset() -> None:
    """Clear the timed-catch inputs for a new calibration."""
    keys = set(BASE_KEYS) | {key for key in st.session_state if key.startswith("orchard_rep_")}
    for key in keys:
        st.session_state.pop(key, None)


def calculate(readings: list[float | None], duration: float | None) -> dict[str, object]:
    """Calculate average catch, variation, and flow rate from three replications."""
    entered = [reading for reading in readings if valid(reading) and reading >= 0]
    complete = len(entered) == len(REPLICATIONS)
    if not complete:
        return {
            "complete": False,
            "average": None,
            "standard_deviation": None,
            "flow_rate": None,
            "deviations": [None] * len(REPLICATIONS),
        }

    average = sum(entered) / len(entered)
    standard_deviation = statistics.stdev(entered) if len(entered) > 1 else 0.0
    deviations = [(reading - average) / average * 100 if average else None for reading in readings]
    flow_rate = average / duration if valid(duration) and duration > 0 else None
    return {
        "complete": True,
        "average": average,
        "standard_deviation": standard_deviation,
        "flow_rate": flow_rate,
        "deviations": deviations,
    }


def build_summary(
    duration: float | None,
    unit: str,
    readings: list[float | None],
    results: dict[str, object],
) -> str:
    """Build the copy-ready Orchard Sprayer result summary."""
    rows = [
        "ORCHARD SPRAYER CALIBRATION",
        "",
        f"Test duration: {show(duration, 1)} seconds",
        f"Catch-volume unit: {unit}",
        "",
        "Timed catch replications:",
    ]
    for replication, reading in zip(REPLICATIONS, readings):
        deviation = results["deviations"][replication - 1]
        status = "CHECK" if valid(deviation) and abs(deviation) > 10 else ("OK" if valid(deviation) else "—")
        rows.append(
            f"  Rep {replication}: {show(reading)} {unit} | "
            f"{show(deviation, 1)}% vs mean | {status}"
        )
    rows.extend(
        [
            "",
            f"Average catch: {show(results['average'])} {unit}",
            f"Standard deviation across reps: {show(results['standard_deviation'])} {unit}",
            f"Average flow rate: {show(results['flow_rate'])} {unit}/s",
        ]
    )
    return "\n".join(rows)


def build_excel_export(
    duration: float | None,
    unit: str,
    readings: list[float | None],
    results: dict[str, object],
) -> bytes:
    """Create an Excel record of the timed-catch test and its current results."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Orchard calibration"

    dark_green = "0B5B3D"
    soft_green = "E8F4ED"
    pale_yellow = "FFF2CC"
    white = "FFFFFF"
    thin_green = Side(style="thin", color="7FA98B")
    section_fill = PatternFill("solid", fgColor=dark_green)
    input_fill = PatternFill("solid", fgColor=pale_yellow)
    result_fill = PatternFill("solid", fgColor=soft_green)
    border = Border(left=thin_green, right=thin_green, top=thin_green, bottom=thin_green)

    def section_header(row: int, title: str) -> None:
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        cell = sheet.cell(row, 1, title)
        cell.font = Font(bold=True, color=white)
        cell.fill = section_fill

    sheet.merge_cells("A1:D1")
    sheet["A1"] = "Orchard Sprayer Calibration Results"
    sheet["A1"].font = Font(bold=True, color=white, size=16)
    sheet["A1"].fill = section_fill
    sheet["A1"].alignment = Alignment(horizontal="center")
    sheet.row_dimensions[1].height = 26

    section_header(3, "Test setup")
    setup = [("Test duration (seconds)", duration), ("Catch-volume unit", unit)]
    for row, (label, value) in enumerate(setup, start=4):
        sheet.cell(row, 1, label)
        input_cell = sheet.cell(row, 2, value)
        input_cell.fill = input_fill
        input_cell.font = Font(color="0000FF")
        input_cell.border = border
        if isinstance(value, (int, float)):
            input_cell.number_format = "0.000"
        sheet.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)

    section_header(7, "Timed catch replications")
    for column, heading in enumerate(["Replication", f"Catch ({unit})", "% deviation from mean", "Status"], start=1):
        cell = sheet.cell(8, column, heading)
        cell.font = Font(bold=True, color=white)
        cell.fill = section_fill
        cell.border = border

    for replication, reading in zip(REPLICATIONS, readings):
        row = 8 + replication
        deviation = results["deviations"][replication - 1]
        status = "Check" if valid(deviation) and abs(deviation) > 10 else ("OK" if valid(deviation) else "")
        values = [replication, reading, deviation / 100 if valid(deviation) else None, status]
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row, column, value)
            cell.border = border
            if column == 2:
                cell.fill = input_fill
                cell.font = Font(color="0000FF")
                cell.number_format = "0.000"
            elif column in (3, 4):
                cell.fill = result_fill
                if column == 3:
                    cell.number_format = "0.0%"

    section_header(13, "Calculated results")
    result_rows = [
        (f"Average catch ({unit})", results["average"], "0.000"),
        (f"Standard deviation ({unit})", results["standard_deviation"], "0.000"),
        (f"Average flow rate ({unit}/s)", results["flow_rate"], "0.000"),
    ]
    for offset, (label, value, number_format) in enumerate(result_rows, start=1):
        row = 13 + offset
        sheet.cell(row, 1, label)
        result_cell = sheet.cell(row, 2, value)
        result_cell.fill = result_fill
        result_cell.font = Font(bold=True)
        result_cell.border = border
        result_cell.number_format = number_format
        sheet.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
        for column in range(1, 5):
            sheet.cell(row, column).border = border

    sheet.merge_cells("A19:D19")
    sheet["A19"] = "Yellow cells record entered inputs; green cells record results calculated by the app at export time."
    sheet["A19"].font = Font(italic=True, color="555555")
    sheet.column_dimensions["A"].width = 33
    sheet.column_dimensions["B"].width = 20
    sheet.column_dimensions["C"].width = 28
    sheet.column_dimensions["D"].width = 14
    sheet.freeze_panes = "A4"

    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


st.title("Orchard Sprayer Calibration")
st.write(
    "Measure the total catch over a fixed time three times. The calculator averages "
    "the catches, checks replication consistency, and calculates flow rate."
)

with st.expander("Field test method", expanded=True):
    st.markdown(
        """1. Choose the catch-volume unit used for the test: **ml** or **L**.
2. Set a fixed test duration, for example **30 seconds**.
3. Run the orchard sprayer for that duration and measure the total catch volume.
4. Repeat the timed catch test three times and enter the three volumes below.

Use clean water. Keep the same sprayer setting and test duration for every replication."""
    )

st.header("1. Timed catch test")
setup_column, unit_column = st.columns(2)
with setup_column:
    duration = st.number_input(
        "Test duration (seconds)",
        min_value=0.1,
        value=30.0,
        step=0.1,
        key="orchard_duration",
    )
with unit_column:
    unit = st.selectbox("Catch-volume unit", options=UNITS, key="orchard_unit")

st.subheader("Catch-volume replications")
st.caption(f"Enter the total volume caught during {show(duration, 1)} seconds for each replication.")
replication_columns = st.columns(3)
readings: list[float | None] = []
for column, replication in zip(replication_columns, REPLICATIONS):
    with column:
        reading = st.number_input(
            f"Rep {replication} catch ({unit})",
            min_value=0.0,
            value=None,
            step=0.1,
            placeholder=f"{unit}",
            key=f"orchard_rep_{replication}",
        )
        readings.append(reading)

results = calculate(readings, duration)
entered = len([reading for reading in readings if valid(reading) and reading >= 0])
st.caption(f"**{entered} of 3** replications entered")

average_metric, standard_deviation_metric, flow_metric = st.columns(3)
average_metric.metric(f"Average catch ({unit})", show(results["average"]))
standard_deviation_metric.metric(f"Standard deviation ({unit})", show(results["standard_deviation"]))
flow_metric.metric(f"Average flow rate ({unit}/s)", show(results["flow_rate"]))

st.subheader("Replication deviation from average")
st.caption("A check is flagged when a replication differs by more than ±10% from the average catch volume.")
if results["complete"]:
    deviation_columns = st.columns(3)
    for column, replication, deviation in zip(deviation_columns, REPLICATIONS, results["deviations"]):
        status = "Check" if abs(deviation) > 10 else "OK"
        column.metric(f"Rep {replication} deviation", f"{deviation:+.1f}%", status)
else:
    st.info("Complete all three catch-volume replications to calculate the average, standard deviation, and flow rate.")

outliers = [replication for replication, deviation in zip(REPLICATIONS, results["deviations"]) if valid(deviation) and abs(deviation) > 10]
if entered < 3:
    remaining = 3 - entered
    st.warning(f"Enter {remaining} remaining replication{'s' if remaining != 1 else ''} to calculate the result.")
if outliers:
    st.warning(f"Check Rep {' and '.join(map(str, outliers))}: catch volume differs by more than 10% from the average.")
st.info(f"**Average flow rate ({unit}/s)** = average catch volume ({unit}) ÷ test duration (seconds).")

summary = build_summary(duration, unit, readings, results)
excel_file = build_excel_export(duration, unit, readings, results)
copy_column, export_column, reset_column = st.columns(3)
with copy_column:
    with st.expander("Copy results"):
        st.caption("Use the copy icon in the top-right of the result box.")
        st.code(summary, language=None)
with export_column:
    export_timestamp = datetime.now(timezone(timedelta(hours=2))).strftime("%Y_%m_%d %H-%M")
    st.download_button(
        "Export results to Excel",
        data=excel_file,
        file_name=f"orchard_sprayer_calibration_{export_timestamp}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        use_container_width=True,
    )
with reset_column:
    st.button("Reset calculator", on_click=reset, type="secondary", use_container_width=True)

st.divider()
st.caption(
    "This is a field-calibration aid. Follow equipment guidance, PPE requirements, "
    "and local operating procedures. Check entries and units before acting on results."
)
