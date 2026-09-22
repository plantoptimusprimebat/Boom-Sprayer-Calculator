"""Boom Sprayer Calibration calculator with three nozzle-catch replications."""
from __future__ import annotations

import math
from datetime import datetime, timedelta, timezone
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
import streamlit as st

st.set_page_config(page_title="Boom Sprayer Calibration", page_icon="🚜", layout="centered")

st.markdown(
    """
<style>
:root { color-scheme: dark; }
html, body, [data-testid="stAppViewContainer"], .stApp {
    background: #101914 !important;
    color: #f3f8f4 !important;
}
[data-testid="stHeader"] { background: #101914 !important; }
[data-testid="stMainBlockContainer"], .stApp p, .stApp li, .stApp label,
[data-testid="stCaptionContainer"] { color: #f3f8f4 !important; }
h1, h2, h3 { color: #8bd5a7 !important; }
[data-testid="stExpander"], [data-testid="stExpanderDetails"] {
    background: #18251d !important; border-color: #385443 !important;
}
[data-testid="stExpander"] summary, [data-testid="stExpander"] summary span {
    color: #f3f8f4 !important;
}
div[data-baseweb="input"] { background: #1a2920 !important; border-color: #52705d !important; }
div[data-baseweb="input"] input { color: #f3f8f4 !important; -webkit-text-fill-color: #f3f8f4 !important; }
div[data-baseweb="input"] input::placeholder { color: #aabaae !important; -webkit-text-fill-color: #aabaae !important; opacity: 1; }
div[data-testid="stMetric"] { background: #1a2920 !important; border: 1px solid #385443; border-radius: 10px; padding: .65rem .8rem; }
div[data-testid="stMetricLabel"] p { color: #b9d1c0 !important; }
div[data-testid="stMetricValue"] { color: #8bd5a7 !important; }
[data-testid="stAlert"], [data-testid="stCode"] { color: #f3f8f4 !important; }
[data-testid="stCode"] { background: #18251d !important; }
</style>
""",
    unsafe_allow_html=True,
)

BASE_KEYS = [
    "distance",
    "travel_time",
    "spray_width",
    "nozzle_count",
    "tank_volume",
    "product_count",
]
UNIT_OPTIONS = ("ml", "g", "L", "kg")
REPLICATIONS = (1, 2, 3)


def valid(value: object) -> bool:
    """Return true only for finite numbers."""
    return value is not None and isinstance(value, (int, float)) and math.isfinite(value)


def show(value: object, decimals: int = 1) -> str:
    """Format a number for display or return a dash when unavailable."""
    return f"{value:,.{decimals}f}" if valid(value) else "—"


def reset() -> None:
    """Restore the original calculator state, including reps and product rows."""
    keys = set(BASE_KEYS) | {
        key
        for key in st.session_state
        if key.startswith("rep_") or key.startswith("product_")
    }
    for key in keys:
        st.session_state.pop(key, None)


def replication_readings(nozzle_count: int) -> list[list[float | None]]:
    """Return nozzle catches, grouped by replication."""
    return [
        [st.session_state.get(f"rep_{replication}_nozzle_{index}") for index in range(nozzle_count)]
        for replication in REPLICATIONS
    ]


def calculation_results(readings_by_rep: list[list[float | None]]) -> dict[str, object]:
    """Calculate calibration values only when all three replications are complete."""
    nozzle_count = len(readings_by_rep[0])
    complete = all(
        len([value for value in replication if valid(value) and value >= 0]) == nozzle_count
        for replication in readings_by_rep
    )
    entered_per_rep = [len([value for value in replication if valid(value) and value >= 0]) for replication in readings_by_rep]
    if not complete:
        return {
            "complete": False,
            "entered_per_rep": entered_per_rep,
            "nozzle_means": [None] * nozzle_count,
            "overall_mean": None,
            "low": None,
            "high": None,
            "average_total_ml": None,
            "deviations": [None] * nozzle_count,
        }

    nozzle_means = [sum(replication[index] for replication in readings_by_rep) / len(REPLICATIONS) for index in range(nozzle_count)]
    overall_mean = sum(nozzle_means) / nozzle_count
    deviations = [
        (nozzle_mean - overall_mean) / overall_mean * 100 if overall_mean else None
        for nozzle_mean in nozzle_means
    ]
    average_total_ml = sum(sum(replication) for replication in readings_by_rep) / len(REPLICATIONS)
    return {
        "complete": True,
        "entered_per_rep": entered_per_rep,
        "nozzle_means": nozzle_means,
        "overall_mean": overall_mean,
        "low": min(nozzle_means),
        "high": max(nozzle_means),
        "average_total_ml": average_total_ml,
        "deviations": deviations,
    }


def product_entries(product_count: int, area_per_tank: float | None) -> list[dict[str, object]]:
    """Render product rows and return their calculated tank quantities."""
    entries: list[dict[str, object]] = []
    for index in range(product_count):
        st.markdown(f"##### Product {index + 1}")
        name_column, rate_column, unit_column, amount_column = st.columns([1.5, 1.15, 0.85, 1.15])
        with name_column:
            name = st.text_input(
                "Product name",
                placeholder=f"Product {index + 1}",
                key=f"product_name_{index}",
            ).strip()
        with rate_column:
            rate = st.number_input(
                "Label rate (per ha)",
                min_value=0.0,
                step=0.1,
                value=None,
                placeholder="e.g. 150",
                key=f"product_rate_{index}",
            )
        with unit_column:
            unit = st.selectbox("Rate unit", options=UNIT_OPTIONS, key=f"product_unit_{index}")
        amount = area_per_tank * rate if valid(area_per_tank) and valid(rate) else None
        with amount_column:
            st.metric(f"Per tank ({unit})", show(amount, 3))
        entries.append({"name": name or f"Product {index + 1}", "rate": rate, "unit": unit, "amount": amount})
    return entries


def build_summary(
    distance: float,
    seconds: float | None,
    width: float | None,
    readings_by_rep: list[list[float | None]],
    results: dict[str, object],
    spray_volume: float | None,
    speed: float | None,
    tank_volume: float | None,
    area_per_tank: float | None,
    products: list[dict[str, object]],
) -> str:
    """Build the copy-ready result summary."""
    rows = [
        "BOOM SPRAYER CALIBRATION — 3 REPLICATIONS",
        "",
        f"Test distance: {show(distance)} m",
        f"Time over distance: {show(seconds)} s",
        f"Spray width: {show(width, 2)} m",
        "",
    ]
    for replication_number, readings in enumerate(readings_by_rep, start=1):
        rows.append(f"Replication {replication_number} nozzle catches:")
        rows.extend(f"  Nozzle {index + 1}: {show(reading)} ml" for index, reading in enumerate(readings))
        rows.append("")

    rows.extend([
        "Average results across 3 replications:",
        f"Average total catch: {show(results['average_total_ml'] / 1000 if valid(results['average_total_ml']) else None, 3)} L",
        f"Mean per nozzle: {show(results['overall_mean'])} ml",
        f"Lowest / highest nozzle mean: {show(results['low'])} / {show(results['high'])} ml",
        f"Travel speed: {show(speed, 2)} km/h",
        f"Spray volume: {show(spray_volume)} L/ha",
        "",
        "Nozzle mean and deviation:",
    ])
    for index, (nozzle_mean, deviation) in enumerate(zip(results["nozzle_means"], results["deviations"]), start=1):
        status = "CHECK" if valid(deviation) and abs(deviation) > 10 else ("OK" if valid(deviation) else "—")
        rows.append(f"  Nozzle {index}: {show(nozzle_mean)} ml | {show(deviation)}% | {status}")
    rows.extend([
        "",
        f"Tank volume: {show(tank_volume, 0)} L",
        f"Area per tank: {show(area_per_tank, 2)} ha",
        "",
        "Products per tank:",
    ])
    for product in products:
        rows.append(
            f"  {product['name']}: {show(product['rate'], 3)} {product['unit']}/ha → "
            f"{show(product['amount'], 3)} {product['unit']} per tank"
        )
    return "\n".join(rows)


def build_excel_export(
    distance: float,
    seconds: float | None,
    width: float | None,
    readings_by_rep: list[list[float | None]],
    results: dict[str, object],
    spray_volume: float | None,
    speed: float | None,
    tank_volume: float | None,
    area_per_tank: float | None,
    products: list[dict[str, object]],
) -> bytes:
    """Create an Excel record containing all three reps and current results."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Calibration"

    dark_green = "0B5B3D"
    soft_green = "E8F4ED"
    pale_yellow = "FFF2CC"
    white = "FFFFFF"
    thin_green = Side(style="thin", color="7FA98B")
    section_fill = PatternFill("solid", fgColor=dark_green)
    input_fill = PatternFill("solid", fgColor=pale_yellow)
    result_fill = PatternFill("solid", fgColor=soft_green)
    border = Border(left=thin_green, right=thin_green, top=thin_green, bottom=thin_green)

    def section_header(row: int, title: str, end_column: int = 7) -> None:
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=end_column)
        cell = sheet.cell(row, 1, title)
        cell.font = Font(bold=True, color=white)
        cell.fill = section_fill

    sheet.merge_cells("A1:G1")
    sheet["A1"] = "Boom Sprayer Calibration Results — 3 Replications"
    sheet["A1"].font = Font(bold=True, color=white, size=16)
    sheet["A1"].fill = section_fill
    sheet["A1"].alignment = Alignment(horizontal="center")
    sheet.row_dimensions[1].height = 26

    section_header(3, "Calibration inputs")
    inputs = [
        ("Test distance (m)", distance),
        ("Time over distance (seconds)", seconds),
        ("Spray width (m)", width),
        ("Number of boom nozzles", len(readings_by_rep[0])),
        ("Tank volume (L)", tank_volume),
    ]
    for row, (label, value) in enumerate(inputs, start=4):
        sheet.cell(row, 1, label)
        value_cell = sheet.cell(row, 2, value)
        value_cell.fill = input_fill
        value_cell.border = border
        value_cell.font = Font(color="0000FF")
        value_cell.number_format = "0.000"
        sheet.merge_cells(start_row=row, start_column=2, end_row=row, end_column=7)

    nozzle_header_row = 10
    section_header(nozzle_header_row, "Nozzle catches and deviation from mean")
    headings = ["Nozzle", "Rep 1 (ml)", "Rep 2 (ml)", "Rep 3 (ml)", "Mean catch (ml)", "% deviation", "Status"]
    for column, heading in enumerate(headings, start=1):
        cell = sheet.cell(nozzle_header_row + 1, column, heading)
        cell.font = Font(bold=True, color=white)
        cell.fill = section_fill
        cell.border = border

    first_nozzle_row = nozzle_header_row + 2
    for index in range(len(readings_by_rep[0])):
        row = first_nozzle_row + index
        sheet.cell(row, 1, index + 1)
        for replication_index, readings in enumerate(readings_by_rep, start=2):
            input_cell = sheet.cell(row, replication_index, readings[index])
            input_cell.fill = input_fill
            input_cell.font = Font(color="0000FF")
            input_cell.number_format = "0.000"
        mean_cell = sheet.cell(row, 5, results["nozzle_means"][index])
        deviation = results["deviations"][index]
        deviation_cell = sheet.cell(row, 6, deviation / 100 if valid(deviation) else None)
        deviation_cell.number_format = "0.0%"
        status = "Check" if valid(deviation) and abs(deviation) > 10 else ("OK" if valid(deviation) else "")
        sheet.cell(row, 7, status)
        for column in range(1, 8):
            cell = sheet.cell(row, column)
            cell.border = border
            if column in (5, 6, 7):
                cell.fill = result_fill
        mean_cell.number_format = "0.000"

    result_header_row = first_nozzle_row + len(readings_by_rep[0]) + 2
    section_header(result_header_row, "Calculated results")
    average_total_litres = results["average_total_ml"] / 1000 if valid(results["average_total_ml"]) else None
    result_rows = [
        ("Average total catch (L)", average_total_litres, "0.000"),
        ("Mean per nozzle (ml)", results["overall_mean"], "0.000"),
        ("Lowest nozzle mean (ml)", results["low"], "0.000"),
        ("Highest nozzle mean (ml)", results["high"], "0.000"),
        ("Spray volume (L/ha)", spray_volume, "0.000"),
        ("Travel speed (km/h)", speed, "0.00"),
        ("Area per tank (ha)", area_per_tank, "0.000"),
    ]
    for offset, (label, value, number_format) in enumerate(result_rows, start=1):
        row = result_header_row + offset
        sheet.cell(row, 1, label)
        result_cell = sheet.cell(row, 2, value)
        result_cell.number_format = number_format
        result_cell.fill = result_fill
        result_cell.font = Font(bold=True)
        for column in range(1, 8):
            sheet.cell(row, column).border = border
        sheet.merge_cells(start_row=row, start_column=2, end_row=row, end_column=7)

    product_header_row = result_header_row + len(result_rows) + 2
    section_header(product_header_row, "Products per tank", end_column=4)
    for column, heading in enumerate(["Product", "Label rate", "Unit", "Required per tank"], start=1):
        cell = sheet.cell(product_header_row + 1, column, heading)
        cell.font = Font(bold=True, color=white)
        cell.fill = section_fill
        cell.border = border

    for index, product in enumerate(products, start=1):
        row = product_header_row + 1 + index
        values = [product["name"], product["rate"], product["unit"], product["amount"]]
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row, column, value)
            cell.border = border
            if column < 4:
                cell.fill = input_fill
                cell.font = Font(color="0000FF")
            else:
                cell.fill = result_fill
                cell.font = Font(bold=True)
                cell.number_format = "0.000"
        sheet.cell(row, 2).number_format = "0.000"

    note_row = product_header_row + len(products) + 4
    sheet.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=7)
    sheet.cell(note_row, 1, "Yellow cells record entered inputs; green cells record results calculated by the app at export time.")
    sheet.cell(note_row, 1).font = Font(italic=True, color="555555")

    for column, width_value in {"A": 28, "B": 16, "C": 16, "D": 16, "E": 19, "F": 15, "G": 14}.items():
        sheet.column_dimensions[column].width = width_value
    sheet.freeze_panes = "A4"

    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


st.title("Boom Sprayer Calibration")
st.write(
    "Enter the catch from every nozzle for three replications. The calculator averages "
    "the replications, identifies uneven nozzles, and calculates spray volume."
)

with st.expander("Field test method", expanded=True):
    st.markdown(
        """1. Measure the test distance: normally **100 m for a tractor** or **30 m for a knapsack**.
2. Travel that distance at normal spraying speed and record the time.
3. Run the sprayer for the same time and catch the output from **every nozzle on the boom**. Repeat the catch test three times and enter each reading below.

Use clean water. Clean or replace blocked, worn, or visibly uneven nozzles before spraying."""
    )

st.header("1. Calibration details")
st.caption("The nozzle count should match the number of nozzles for the boom width entered. The same distance, time, and width apply to all three replications.")
column_1, column_2 = st.columns(2)
with column_1:
    distance = st.number_input("Test distance (m)", min_value=0.1, value=100.0, step=0.1, key="distance")
    width = st.number_input("Spray width (m)", min_value=0.0, value=None, step=0.01, placeholder="e.g. 12", key="spray_width")
with column_2:
    seconds = st.number_input("Time over distance (seconds)", min_value=0.0, value=None, step=0.1, placeholder="e.g. 60", key="travel_time")
    count = int(st.number_input("Number of boom nozzles", min_value=1, max_value=120, value=4, step=1, key="nozzle_count"))

st.subheader("Individual nozzle catches")
st.caption("Enter catch volume in ml for the same travel time. Complete all three replications before results are calculated.")
for replication in REPLICATIONS:
    st.markdown(f"#### Rep {replication}")
    for start in range(0, count, 3):
        columns = st.columns(3)
        for column, index in zip(columns, range(start, min(start + 3, count))):
            with column:
                st.number_input(
                    f"Nozzle {index + 1} (ml)",
                    min_value=0.0,
                    value=None,
                    step=1.0,
                    placeholder="ml",
                    key=f"rep_{replication}_nozzle_{index}",
                )

readings_by_rep = replication_readings(count)
results = calculation_results(readings_by_rep)
average_total_litres = results["average_total_ml"] / 1000 if valid(results["average_total_ml"]) else None
spray_volume = 10_000 * average_total_litres / (distance * width) if valid(average_total_litres) and valid(width) and width > 0 else None
speed = distance / seconds * 3.6 if valid(seconds) and seconds > 0 else None

entered_caption = " • ".join(f"Rep {rep}: **{entered} of {count}**" for rep, entered in zip(REPLICATIONS, results["entered_per_rep"]))
st.caption(entered_caption + " readings entered")
metric_1, metric_2, metric_3, metric_4 = st.columns(4)
metric_1.metric("Average total catch (L)", show(average_total_litres, 3))
metric_2.metric("Mean per nozzle (ml)", show(results["overall_mean"]))
metric_3.metric("Lowest / highest mean (ml)", f"{show(results['low'], 0)} / {show(results['high'], 0)}" if results["complete"] else "—")
metric_4.metric("Spray volume (L/ha)", show(spray_volume))

st.subheader("Nozzle deviation from mean")
st.caption("Each nozzle result is the average of its three replications. A check is flagged when a nozzle differs by more than ±10% from the overall mean.")
if results["complete"]:
    # The card count follows the nozzle count selected above. Four cards are
    # shown per row so the result remains readable on smaller screens.
    for start in range(0, count, 4):
        indices = range(start, min(start + 4, count))
        deviation_columns = st.columns(len(indices))
        for column, index in zip(deviation_columns, indices):
            deviation = results["deviations"][index]
            status = "Check" if abs(deviation) > 10 else "OK"
            column.metric(f"Nozzle {index + 1} deviation", f"{deviation:+.1f}%", status)
else:
    st.info("Complete all nozzle readings in Rep 1, Rep 2, and Rep 3 to show the percentage deviation for every nozzle.")

outliers = [index + 1 for index, deviation in enumerate(results["deviations"]) if valid(deviation) and abs(deviation) > 10]
for replication, entered in zip(REPLICATIONS, results["entered_per_rep"]):
    remaining = count - entered
    if remaining:
        st.warning(f"Rep {replication}: enter {remaining} remaining nozzle reading{'s' if remaining != 1 else ''} to calculate the result.")
if outliers:
    st.warning(f"Check nozzle{'s' if len(outliers) > 1 else ''} {', '.join(map(str, outliers))}: average output differs by more than 10% from the overall mean.")
if valid(speed) and not 3 <= speed <= 12:
    st.warning("Travel speed is outside the common 3–12 km/h range; verify the measurement and safe operating speed.")
st.info("**Spray volume (L/ha)** = 10,000 × average total nozzle catch across 3 reps (L) ÷ [test distance (m) × spray width (m)].")

st.header("2. Speed and products per tank")
st.caption("Enter the tank volume once, then add each product and its registered label rate.")
tank_column, product_count_column = st.columns(2)
with tank_column:
    tank_volume = st.number_input("Tank volume (L)", min_value=0.0, value=None, step=1.0, placeholder="e.g. 600", key="tank_volume")
with product_count_column:
    product_count = int(st.number_input("Number of products to add", min_value=1, max_value=10, value=1, step=1, key="product_count"))

area_per_tank = tank_volume / spray_volume if valid(tank_volume) and valid(spray_volume) and spray_volume > 0 else None
speed_metric, area_metric = st.columns(2)
speed_metric.metric("Travel speed (km/h)", show(speed, 2))
area_metric.metric("Area per tank (ha)", show(area_per_tank, 2))

st.subheader("Products to add")
st.caption("Use the product's registered label rate. Product amounts are calculated from the area covered by one tank.")
products = product_entries(product_count, area_per_tank)
st.info("**Product per tank** = area per tank (ha) × label rate (unit/ha). Check each product label and approved tank-mix guidance before application.")

st.header("Adjusting spray volume")
st.table(
    {
        "To increase spray volume": ["Larger nozzle", "Higher pressure", "Lower speed"],
        "To reduce spray volume": ["Smaller nozzle", "Lower pressure", "Higher speed"],
    }
)

summary = build_summary(distance, seconds, width, readings_by_rep, results, spray_volume, speed, tank_volume, area_per_tank, products)
excel_file = build_excel_export(distance, seconds, width, readings_by_rep, results, spray_volume, speed, tank_volume, area_per_tank, products)
copy_column, export_column, reset_column = st.columns(3)
with copy_column:
    with st.expander("Copy results"):
        st.caption("Use the copy icon in the top-right of the result box.")
        st.code(summary, language=None)
with export_column:
    export_timestamp = datetime.now(timezone(timedelta(hours=2))).strftime("%Y_%m_%d %H-%M")
    st.download_button(
        "Export results to Excel",
        data=excel_file,
        file_name=f"boom_sprayer_calibration_{export_timestamp}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        use_container_width=True,
    )
with reset_column:
    st.button("Reset calculator", on_click=reset, type="secondary", use_container_width=True)

st.divider()
st.caption(
    "This is a field-calibration aid. Follow the registered product label, equipment guidance, "
    "PPE requirements, and local operating procedures. Check entries and units before mixing or application."
)
